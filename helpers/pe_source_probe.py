"""Read-only PE source retrieval/parse probe. Does not publish or refresh app data."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO, StringIO
import json
import os
from pathlib import Path
import re
import sys
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen

from bs4 import BeautifulSoup
import openpyxl
import pandas as pd
from pypdf import PdfReader

SEC_INDEX = "https://www.sec.gov/data-research/investment-management-data"
SEC_FALLBACK = "https://www.sec.gov/files/investment/private-funds-statistics-2025-q4-supporting-data.xlsx"
KPMG_INDEX = "https://kpmg.com/xx/en/media/press-releases.html"
KPMG_RELEASES = (
    "https://kpmg.com/xx/en/media/press-releases/2025/07/global-pe-investment-slows-sharply-amid-rising-geopolitical-and-trade-uncertainty.html",
    "https://kpmg.com/xx/en/media/press-releases/2026/07/global-pe-market-shows-resilience-with-over-1-trillion-in-pe-investment-in-first-half-of-2026.html",
)
ACG_INDEX = "https://www.acg.org/news-trends/news"
ACG_RELEASES = (
    "https://www.acg.org/news-trends/news/gf-data-reports-show-stronger-financing-markets-fueled-ma-gains-q1",
    "https://www.acg.org/news-trends/news/gf-data-reports-show-steady-middle-market-deal-flow-amid-more-selective",
)
CALPERS_URL = "https://www.calpers.ca.gov/investments/about-investment-office/investment-organization/pep-fund-performance"
WSIB_INDEX = "https://www.sib.wa.gov/reports.html"
OREGON_URL = "https://www.oregon.gov/treasury/invested-for-oregon/Documents/Invested-for-OR-47OIC-Agenda-and-Minutes/2026/Regular-Meeting-January-21-2026.pdf"
SOURCES = ("sec", "kpmg", "acg", "calpers", "wsib", "oregon")
MAX_BYTES = 12_000_000
TIMEOUT = 25


class ProbeError(ValueError):
    """Expected source, download, or parsing failure."""


def _host_allowed(url: str, allowed_hosts: tuple[str, ...]) -> bool:
    parsed = urlparse(url)
    return parsed.scheme == "https" and (parsed.hostname or "").lower() in allowed_hosts


def fetch(url: str, hosts: tuple[str, ...], *, max_bytes=MAX_BYTES) -> tuple[bytes, dict]:
    if not _host_allowed(url, hosts):
        raise ProbeError(f"Unapproved source URL: {url}")
    agent = (os.environ.get("SEC_USER_AGENT", "").strip() if "sec.gov" in urlparse(url).hostname else "")
    agent = agent or "AI-Macro-PE-Source-Probe/1.0 (public-market research)"
    request = Request(url, headers={"User-Agent": agent, "Accept": "text/html,application/pdf,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*"})
    try:
        with urlopen(request, timeout=TIMEOUT) as response:
            final_url = response.geturl()
            if not _host_allowed(final_url, hosts):
                raise ProbeError(f"Unexpected redirect host: {final_url}")
            data = response.read(max_bytes + 1)
            content_type = response.headers.get("Content-Type", "")
            status = response.status
    except (HTTPError, URLError, TimeoutError, OSError) as exc:
        raise ProbeError(f"{type(exc).__name__}: {exc}") from exc
    if len(data) > max_bytes:
        raise ProbeError(f"Response exceeds {max_bytes:,} bytes")
    if not data:
        raise ProbeError("Empty response")
    return data, {"url": final_url, "http_status": status, "bytes": len(data), "content_type": content_type}


def html_text(data: bytes) -> str:
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    return " ".join(soup.get_text(" ", strip=True).split())


def require_html(data: bytes, *, keyword: str) -> str:
    if b"<html" not in data[:5000].lower() and b"<!doctype html" not in data[:5000].lower():
        raise ProbeError("Expected HTML, got another response type")
    text = html_text(data)
    if keyword.casefold() not in text.casefold() or len(text) < 200:
        raise ProbeError(f"HTML does not contain expected subject: {keyword}")
    return text


def same_domain_links(data: bytes, base: str, hosts: tuple[str, ...]):
    soup = BeautifulSoup(data, "html.parser")
    for tag in soup.select("a[href]"):
        url = urljoin(base, tag.get("href", ""))
        if _host_allowed(url, hosts):
            yield " ".join(tag.get_text(" ", strip=True).split()), url


def parse_sec_workbook(data: bytes) -> dict:
    if not data.startswith(b"PK\x03\x04"):
        raise ProbeError("SEC workbook is not an XLSX/ZIP payload")
    try:
        book = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (OSError, ValueError, ImportError) as exc:
        raise ProbeError(f"XLSX cannot be parsed: {exc}") from exc
    try:
        hits, periods, numeric_cells = [], set(), 0
        for sheet in book.worksheets:
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 800:
                    break
                cells = [str(x).strip() if x is not None else "" for x in row]
                if index < 30:
                    for cell in cells:
                        periods.update(re.findall(r"20\d{2}Q[1-4]", cell, re.I))
                if any(re.search(r"private\s+equity\s+fund", cell, re.I) for cell in cells):
                    numeric = [x for x in row if isinstance(x, (int, float)) and not isinstance(x, bool) and x >= 0]
                    if numeric:
                        numeric_cells += len(numeric)
                        hits.append(sheet.title)
        if len(set(hits)) < 2 or numeric_cells < 4:
            raise ProbeError(f"Workbook opens but PE quantitative table detection failed ({len(set(hits))} sheets)")
        return {"sheets": len(book.sheetnames), "pe_sheets": sorted(set(hits))[:16], "quantitative_cells": numeric_cells,
                "historical_periods_detected": sorted(periods)[-12:]}
    finally:
        book.close()


def parse_kpmg_release(data: bytes) -> dict:
    text = require_html(data, keyword="private equity")
    amounts = re.findall(r"\$\s*\d[\d,.]*\s*(?:billion|trillion)", text, re.I)
    if len(amounts) < 4 or not re.search(r"\bdeals?\b", text, re.I) or not re.search(r"\bexit", text, re.I) or not re.search(r"\bfundrais", text, re.I):
        raise ProbeError("KPMG release lacks extractable quantitative deal, exit and fundraising context")
    return {"currency_amounts": len(amounts), "amount_examples": amounts[:6],
            "has_deals": True, "has_exits": True, "has_fundraising": True}


def parse_acg_release(data: bytes) -> dict:
    text = require_html(data, keyword="GF Data")
    if not re.search(r"\bEBITDA\b", text, re.I):
        raise ProbeError("No EBITDA valuation basis")
    multiples = re.findall(r"(?<!\w)\d+(?:\.\d+)?x(?!\w)", text, re.I)
    counts = re.findall(r"\b\d{2,4}\s+(?:completed\s+)?transactions\b", text, re.I)
    if not multiples or not counts:
        raise ProbeError("No purchase multiples and transaction counts identified")
    return {"multiples_detected": multiples[:8], "transaction_count_examples": counts[:4]}


def parse_calpers(data: bytes) -> dict:
    text = require_html(data, keyword="Private Equity")
    date = re.search(r"table below reflects.{0,260}?as of\s+([A-Za-z]+\s+\d{1,2},\s*\d{4})", text, re.I)
    if not date:
        raise ProbeError("CalPERS fund table observation date not found")
    try:
        tables = pd.read_html(StringIO(data.decode("utf-8", "replace")), keep_default_na=False)
    except (ValueError, ImportError) as exc:
        raise ProbeError(f"CalPERS tables unreadable: {exc}") from exc
    expected = {"fund", "vintageyear", "cashin", "cashout", "cashoutremainingvalue"}
    matches = []
    for table in tables:
        headings = {re.sub(r"[^a-z0-9]", "", str(col).lower()) for col in table.columns}
        if expected <= headings:
            matches.append(table)
    if len(matches) != 1 or len(matches[0]) < 100:
        raise ProbeError(f"No sufficiently broad CalPERS investment table ({len(matches)} matches)")
    return {"as_of": datetime.strptime(date.group(1), "%B %d, %Y").date().isoformat(),
            "fund_rows": len(matches[0]), "contains_cash_in_out": True,
            "coverage": "active CalPERS PE fund positions, not all industry funds"}


def parse_wsib_index(data: bytes) -> tuple[str, str]:
    require_html(data, keyword="Reports")
    urls = []
    for _, url in same_domain_links(data, WSIB_INDEX, ("www.sib.wa.gov", "sib.wa.gov")):
        match = re.search(r"/docs/reports/quarterly/ir(\d{6})\.pdf(?:\?|$)", url, re.I)
        if match:
            try:
                dt = datetime.strptime(match.group(1), "%m%d%y").date()
                if dt <= datetime.now(timezone.utc).date():
                    urls.append((dt, url))
            except ValueError:
                continue
    if not urls:
        raise ProbeError("No dated PE IRR PDF discovered from WSIB reports index")
    dt, url = max(urls)
    return url, dt.isoformat()


def pdf_text(data: bytes, *, min_pages: int = 2) -> tuple[str, int]:
    if not data.startswith(b"%PDF"):
        raise ProbeError("Not a PDF payload")
    try:
        reader = PdfReader(BytesIO(data), strict=False)
        if len(reader.pages) < min_pages or len(reader.pages) > 250:
            raise ProbeError(f"Unusual PDF page count: {len(reader.pages)}")
        return "\n".join((page.extract_text() or "") for page in reader.pages), len(reader.pages)
    except (ValueError, OSError) as exc:
        raise ProbeError(f"PDF parse failed: {exc}") from exc


def parse_wsib_pdf(data: bytes) -> dict:
    text, pages = pdf_text(data)
    dates = re.findall(r"\b\d{1,2}/\d{1,2}/20\d{2}\b", text)
    if pages > 30 or len(dates) < 25:
        raise ProbeError(f"WSIB PDF content incomplete: pages={pages}, fund dates={len(dates)}")
    return {"pages": pages, "rows_with_dates": len(dates), "type": "fund-level LP records"}


def parse_oregon_pdf(data: bytes) -> dict:
    text, pages = pdf_text(data, min_pages=5)
    if "private equity" not in text.casefold() or "cash flow" not in text.casefold():
        raise ProbeError("Oregon agenda PDF lacks PE cash-flow discussion")
    # A chart image is insufficient for a reproducible numeric series. This
    # source is deliberately not counted as structured data until a table exists.
    return {"pages": pages, "pe_cash_flow_text": True, "numeric_series": False,
            "limitation": "Cash-flow observations appear as chart graphics; no validated table parser"}


def discover_press_links(data: bytes, index_url: str, hosts: tuple[str, ...], name: str) -> list[str]:
    require_html(data, keyword="press" if name == "kpmg" else "news")
    matches = []
    for label, url in same_domain_links(data, index_url, hosts):
        searchable = label + " " + url
        if name == "kpmg":
            ok = re.search(r"(?:private.equity|(?:\bpe\b.{0,35}market))", searchable, re.I) and "/media/press-releases/20" in url
        else:
            ok = re.search(r"gf.data", searchable, re.I) and "/news-trends/news/" in url
        if ok:
            matches.append(url)
    return list(dict.fromkeys(matches))[:12]


def report_row(name: str, question: str, links: list[str], *, observation=None, discovered=None,
               notes=None, errors=None, structured=False) -> dict:
    observation = observation or {}
    discovered = discovered or []
    errors = errors or []
    if not observation:
        status = "failed"
    elif observation.get("numeric_series") is False:
        status = "retrievable_only"
    elif errors:
        status = "partial"
    else:
        status = "structured" if structured else "content_verified"
    return {
        "source": name, "question": question,
        "status": status,
        "structured_observations": bool(observation and structured),
        "production_feed_ready": False,
        "source_urls": links, "index_discovery": bool(discovered), "discovered_urls": discovered[:5],
        "extraction": observation, "notes": notes or "", "errors": errors,
    }


def probe_sec() -> dict:
    links, errors, discovered, details = [], [], [], {}
    workbook_url = SEC_FALLBACK
    try:
        body, info = fetch(SEC_INDEX, ("www.sec.gov", "sec.gov"))
        links.append(info["url"])
        require_html(body, keyword="Private Fund Statistics")
        for label, url in same_domain_links(body, SEC_INDEX, ("www.sec.gov", "sec.gov")):
            if "Private Fund Statistics Supporting Data" in label and url.lower().split("?")[0].endswith(".xlsx"):
                workbook_url = url
                discovered.append(url)
                break
        if not discovered:
            errors.append("Latest supporting workbook not discovered; checking a pinned 2025Q4 fallback")
    except ProbeError as exc:
        errors.append(f"Index: {exc}")
    try:
        body, info = fetch(workbook_url, ("www.sec.gov", "sec.gov"))
        links.append(info["url"])
        details = {**parse_sec_workbook(body), "download": info}
    except ProbeError as exc:
        errors.append(f"XLSX: {exc}")
    return report_row("sec", "Industry scale / leverage", links, observation=details, discovered=discovered, errors=errors,
                      structured=bool(details), notes="Form PF filing cycles differ by fund type; a workbook download alone is not a new economic reading.")


def probe_press(name: str) -> dict:
    if name == "kpmg":
        index, urls, hosts = KPMG_INDEX, KPMG_RELEASES, ("kpmg.com",)
        parser, question = parse_kpmg_release, "Deals / exits / fundraising"
    else:
        index, urls, hosts = ACG_INDEX, ACG_RELEASES, ("www.acg.org", "acg.org")
        parser, question = parse_acg_release, "Middle-market entry pricing / leverage"
    links, errors, discovered, periods = [], [], [], []
    try:
        body, info = fetch(index, hosts)
        links.append(info["url"])
        discovered = discover_press_links(body, index, hosts, name)
        if not discovered:
            errors.append("Index accessible but no PE release links discovered")
    except ProbeError as exc:
        errors.append(f"Index: {exc}")
    for url in urls:
        try:
            body, info = fetch(url, hosts)
            links.append(info["url"])
            periods.append({"url": info["url"], "data": parser(body), "bytes": info["bytes"]})
        except ProbeError as exc:
            errors.append(f"Release {url}: {exc}")
    observations = {"release_count": len(periods), "releases": periods} if periods else {}
    if len(periods) != len(urls):
        errors.append(f"Only {len(periods)} of {len(urls)} reference releases parsed")
    return report_row(name, question, links, observation=observations, discovered=discovered, errors=errors,
                      notes="Release prose must be parsed by definition and period before it can populate a production time series.")


def probe_calpers() -> dict:
    try:
        data, info = fetch(CALPERS_URL, ("www.calpers.ca.gov", "calpers.ca.gov"))
        return report_row("calpers", "LP cash in/out", [info["url"]], observation={**parse_calpers(data), "download": info},
                          discovered=[info["url"]], structured=True,
                          notes="Cumulative active positions only; not quarter-specific cash flows or an industry census.")
    except ProbeError as exc:
        return report_row("calpers", "LP cash in/out", [CALPERS_URL], errors=[str(exc)])


def probe_wsib() -> dict:
    links, errors, discovered, details = [], [], [], {}
    try:
        data, info = fetch(WSIB_INDEX, ("www.sib.wa.gov", "sib.wa.gov"))
        links.append(info["url"])
        url, as_of = parse_wsib_index(data)
        discovered.append(url)
        pdf, pdf_info = fetch(url, ("www.sib.wa.gov", "sib.wa.gov"))
        links.append(pdf_info["url"])
        details = {**parse_wsib_pdf(pdf), "as_of": as_of, "download": pdf_info}
    except ProbeError as exc:
        errors.append(str(exc))
    return report_row("wsib", "LP fund returns / possible cash-flow changes", links, observation=details,
                      discovered=discovered, errors=errors,
                      notes="Fund-list coverage and treatment of commitments must be reconciled before differencing quarters.")


def probe_oregon() -> dict:
    try:
        data, info = fetch(OREGON_URL, ("www.oregon.gov", "oregon.gov"))
        details = {**parse_oregon_pdf(data), "download": info}
        return report_row("oregon", "LP cash-flow history", [info["url"]], observation=details,
                          notes="Pinned public presentation; no stable numerical feed established and no automatic new-release discovery.")
    except ProbeError as exc:
        return report_row("oregon", "LP cash-flow history", [OREGON_URL], errors=[str(exc)])


PROBES = {"sec": probe_sec, "kpmg": lambda: probe_press("kpmg"),
          "acg": lambda: probe_press("acg"), "calpers": probe_calpers,
          "wsib": probe_wsib, "oregon": probe_oregon}


def run(selected: list[str]) -> dict:
    rows = []
    for name in selected:
        print(f"Probing {name}...", flush=True)
        try:
            row = PROBES[name]()
        except Exception as exc:
            row = report_row(name, "probe", [], errors=[f"Unexpected {type(exc).__name__}: {exc}"])
        rows.append(row)
        print(f"  {row['status']} | structured={row['structured_observations']} | discovery={row['index_discovery']} | errors={len(row['errors'])}", flush=True)
    return {"run_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "purpose": "Read-only retrieval diagnostics; no dashboard data updated",
            "production_status": "No source certified for production until repeated unattended runs and definition-matched parsing",
            "results": rows,
            "summary": {"sources_tested": len(rows),
                        "validated_content_without_error": sum(x["status"] in {"structured", "content_verified"} for x in rows),
                        "with_structured_observations": sum(x["structured_observations"] for x in rows),
                        "with_index_discovery": sum(x["index_discovery"] for x in rows)}}


def main(argv=None) -> int:
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--output", type=Path, default=Path("pe_probe_output/source_probe.json"))
    args.add_argument("--only", default=",".join(SOURCES), help="Comma-separated source IDs")
    opts = args.parse_args(argv)
    selected = [x.strip().lower() for x in opts.only.split(",") if x.strip()]
    if not selected or len(selected) != len(set(selected)) or any(x not in PROBES for x in selected):
        args.error(f"--only must contain unique source IDs from: {', '.join(SOURCES)}")
    result = run(selected)
    opts.output.parent.mkdir(parents=True, exist_ok=True)
    opts.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Saved {opts.output} | summary: {result['summary']}")
    # A source-specific failure is a useful finding, not an execution failure.
    return 0


if __name__ == "__main__":
    sys.exit(main())
